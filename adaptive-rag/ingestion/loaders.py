"""
Document loaders.

Deliberately built directly on `pypdf` / `python-docx` instead of
`langchain_community` document loaders: the community loader package has a
huge, fast-churning dependency surface for functionality we don't need,
and per-format extraction logic here is simple enough to own directly and
extend for new formats later (see `LOADER_REGISTRY`).

Every loader returns a list of `RawPage` — plain extracted text plus a
1-based page number (documents with no real pagination, e.g. .txt/.md,
are treated as a single "page").
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Protocol


class UnsupportedFileTypeError(ValueError):
    pass


class DocumentLoadError(RuntimeError):
    """Raised for corrupted/unreadable files instead of letting a raw
    library exception escape into the ingestion pipeline."""


@dataclass
class RawPage:
    page_number: int
    text: str


class Loader(Protocol):
    def __call__(self, path: Path) -> list[RawPage]: ...


def load_pdf(path: Path) -> list[RawPage]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except (PdfReadError, OSError, ValueError) as exc:
        raise DocumentLoadError(f"Could not open PDF '{path.name}': {exc}") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentLoadError(
                f"PDF '{path.name}' is password-protected and could not be opened."
            ) from exc

    pages: list[RawPage] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        pages.append(RawPage(page_number=i, text=text))

    if not any(p.text.strip() for p in pages):
        raise DocumentLoadError(
            f"No extractable text found in '{path.name}'. It may be a scanned "
            "image PDF that needs OCR, which is not currently supported."
        )
    return pages


def load_txt(path: Path) -> list[RawPage]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise DocumentLoadError(f"Could not read text file '{path.name}': {exc}") from exc
    if not text.strip():
        raise DocumentLoadError(f"'{path.name}' is empty.")
    return [RawPage(page_number=1, text=text)]


def load_markdown(path: Path) -> list[RawPage]:
    return load_txt(path)


def load_docx(path: Path) -> list[RawPage]:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(str(path))
    except (PackageNotFoundError, OSError, KeyError) as exc:
        raise DocumentLoadError(f"Could not open DOCX '{path.name}': {exc}") from exc

    parts: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise DocumentLoadError(f"No extractable text found in '{path.name}'.")
    # DOCX has no reliable page boundaries without a rendering engine, so
    # the whole document is treated as a single logical page.
    return [RawPage(page_number=1, text=text)]


LOADER_REGISTRY: dict[str, Callable[[Path], list[RawPage]]] = {
    ".pdf": load_pdf,
    ".txt": load_txt,
    ".md": load_markdown,
    ".docx": load_docx,
}


def load_document(path: Path) -> list[RawPage]:
    ext = path.suffix.lower()
    loader = LOADER_REGISTRY.get(ext)
    if loader is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Supported: {sorted(LOADER_REGISTRY)}"
        )
    return loader(path)
